"""Agent 3 — Rédacteur : adapte le CV et la lettre de motivation à une offre.

CV : la forme ne change JAMAIS. Seul le titre du CV (ex. « Chargé de Conformité - Sécurité Financière-Sanctions
Internationales ») est remplacé par un titre adapté à l'offre ; police, couleurs, cadre, colonnes et tout le reste du
texte restent identiques. Le fichier est nommé d'après l'offre.

Lettre : la lettre modèle est gardée telle quelle, avec en plus :
- le bloc destinataire (nom et adresse de l'employeur) aligné à droite ;
- la date du jour, alignée à droite ;
- l'objet et les paragraphes de motivation réécrits pour l'offre (même nombre de paragraphes que le modèle).

Aucun fait n'est inventé (pas de nouvel employeur, diplôme, date ou chiffre). Export Word (.docx) + PDF.
"""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from datetime import datetime

from docx import Document

from services.documents import (SourceDoc, _iter_paragraphs, adapt_letter_docx, letter_structure,
                                read_document, replace_text_everywhere, safe_filename, save_output,
                                standard_letter_docx, to_pdf)
from services.keywords import SOFT, find_skills, guess_offer_fields, match_report

from .base import OUTPUTS, Agent, AgentInfo, ensure_dirs

MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
ROLE_RX = re.compile(r"analyst|charg|consultant|officer|gestionnaire|ingénieur|ingenieur|manager|chef|responsable|assistant|"
                     r"développeur|developpeur|auditeur|contrôleur|controleur|specialist|spécialiste|juriste|comptable|"
                     r"conformit|compliance|kyc|lcb|sanction|risque|crédit|credit", re.I)
NUM_RX = re.compile(r"\d+(?:[.,]\d+)?")
HF_RX = re.compile(r"\s*\(?\b[hfm]\s*/\s*[hfm]\b\)?|\s*\((?:e|é)\)", re.I)
CONTRACT_RX = re.compile(r"\s*[-–|,]?\s*\b(cdi|cdd|int[ée]rim|stage|alternance|freelance)\b.*$", re.I)
# Thèmes affichables dans un titre de CV (repris du vocabulaire de l'offre)
THEMES = {"Sanctions & embargos": "Sanctions Internationales", "LCB-FT": "LCB-FT", "AML / CFT": "AML/CFT", "KYC": "KYC",
          "KYB": "KYB", "Filtrage / screening": "Filtrage", "Surveillance transactionnelle": "Surveillance Transactionnelle",
          "Déclaration de soupçon (Tracfin)": "Déclarations de Soupçon", "Due diligence / EDD": "Due Diligence",
          "PPE / PEP": "PPE", "Fraude": "Lutte contre la Fraude", "Correspondance bancaire": "Correspondance Bancaire",
          "Contrôle permanent": "Contrôle Permanent", "Risque de crédit": "Risque de Crédit",
          "Analyse financière": "Analyse Financière", "Conformité réglementaire": "Conformité Réglementaire",
          "Anti-corruption (Sapin II)": "Anti-Corruption", "Contrôle des exportations": "Contrôle des Exportations",
          "Business analysis": "Business Analysis", "Paiements / PSP": "Paiements", "Crypto-actifs / MiCA": "Crypto-Actifs"}

SYSTEM = (
    "Tu es un expert du recrutement en France (banque, finance, conformité) et un rédacteur de candidatures. Règle "
    "absolue : tu n'inventes RIEN — aucune expérience, aucun employeur, diplôme, certification, date ou chiffre absent du "
    "CV. Français professionnel, précis, sans formules creuses ni superlatifs.")


def today_fr() -> str:
    d = datetime.now()
    return f"{d.day if d.day > 1 else '1er'} {MOIS[d.month - 1]} {d.year}"


def de(x: str) -> str:
    """« de Analyste » → « d'Analyste »."""
    x = str(x or "").strip()
    return ("d'" if re.match(r"[aeiouyhâàéèêëîïôöûüAEIOUYHÂÀÉÈÊËÎÏÔÖÛÜ]", x) else "de ") + x


def clean_title(t: str) -> str:
    t = HF_RX.sub("", str(t or ""))
    t = CONTRACT_RX.sub("", t)
    return re.sub(r"\s{2,}", " ", t).strip(" -–|·,")


@dataclass
class Result:
    kind: str
    docx: bytes
    pdf: bytes
    pdf_method: str
    text: str
    filename: str = "document"
    changes: list = field(default_factory=list)  # (avant, après, raison)
    warnings: list = field(default_factory=list)
    score_before: int = 0
    score_after: int = 0
    missing: list = field(default_factory=list)
    paths: tuple = ("", "")


class AgentRedacteur(Agent):
    info = AgentInfo(
        key="redacteur", nom="Agent Rédacteur CV & lettre", icone="✍️",
        mission="Garde votre CV tel quel et n'adapte que son titre à l'offre ; complète votre lettre avec l'adresse du "
                "destinataire, la date du jour et une motivation propre à l'offre. Export Word et PDF.",
        entrees=["Texte de l'offre (copier-coller)", "CV modèle Word", "Lettre modèle (facultative)"],
        sorties=["CV au titre adapté (.docx + .pdf)", "Lettre adaptée (.docx + .pdf)", "Fichiers nommés d'après l'offre"])

    # ------------------------------------------------------------------ analyse de l'offre
    def analyze_offer(self, text: str, use_ai: bool = True) -> dict:
        base = guess_offer_fields(text)
        base.update({"missions": [], "soft_skills": [], "resume": "", "secteur": ""})
        if use_ai and self.llm.ok:
            prompt = ("Analyse cette offre d'emploi. Réponds en JSON avec les clés : poste (intitulé exact sans H/F ni type de "
                      "contrat), entreprise, lieu, contrat, experience, secteur (activité de l'employeur), missions (liste "
                      "courte), competences_requises (liste), soft_skills (liste), mots_cles (10 à 15 termes), resume "
                      "(2 phrases).\n\nOFFRE :\n" + text[:12000])
            try:
                ai = self.llm.complete(SYSTEM, prompt, json_mode=True, max_tokens=2000)
                for k, v in (ai or {}).items():
                    if v not in (None, "", []):
                        base[k] = v
            except Exception:  # noqa: BLE001 — repli sur l'analyse sans IA
                pass
        base["poste"] = clean_title(base.get("poste"))
        return base

    # ------------------------------------------------------------------ CV : titre uniquement
    @staticmethod
    def _bold_flags(cv: SourceDoc) -> dict[int, bool]:
        flags = {}
        try:
            doc = Document(io.BytesIO(cv.data))
            for i, p in enumerate(_iter_paragraphs(doc)):
                runs = [r for r in p.runs if r.text.strip()]
                flags[i] = bool(runs) and all(bool(r.bold) or bool(r.style and r.style.font.bold) for r in runs)
        except Exception:  # noqa: BLE001
            pass
        return flags

    def title_candidates(self, cv: SourceDoc) -> list[str]:
        """Lignes du CV pouvant être le titre, la plus probable en premier."""
        if not cv.is_docx:
            return []
        bold = self._bold_flags(cv)
        scored = []
        for rank, p in enumerate([p for p in cv.paragraphs if p.text.strip()][:25]):
            t = p.text.strip()
            if not p.editable or len(t) > 120 or len(t.split()) < 2 or re.search(r"[|@]|\d{2}[ .]\d{2}", t):
                continue
            s = 0
            s += 4 if ROLE_RX.search(t) else 0
            s += 2 if bold.get(p.idx) else 0
            s += 1 if re.search(r"\s[-–]\s|-", t) else 0
            s -= rank * 0.1
            if s >= 4:
                scored.append((s, t))
        saved = (self.cfg.user or {}).get("titre_cv", "")
        out = [t for _, t in sorted(scored, key=lambda x: -x[0])]
        if saved and saved in cv.text:
            out = [saved] + [t for t in out if t != saved]
        return list(dict.fromkeys(out))

    def propose_title(self, old_title: str, offer: dict, cv_text: str, use_ai: bool = True, offer_text: str = "") -> str:
        poste = clean_title(offer.get("poste")) or "Poste"
        words = set(re.findall(r"\w{4,}", poste.lower()))
        cands = (["Sécurité Financière"] if re.search(r"s[ée]curit[ée] financi[èe]re", offer_text, re.I) else []) + \
            [THEMES[s] for s in offer.get("competences_requises") or [] if s in THEMES]
        themes = [t for t in dict.fromkeys(cands) if not set(re.findall(r"\w{4,}", t.lower())) & words]
        rule = poste
        if re.search(r"\s[-–]\s", old_title) and themes:
            rule = f"{poste} - {'-'.join(dict.fromkeys(themes[:2]))}"
        if len(rule) > max(len(old_title) + 15, 70):
            rule = poste
        if not (use_ai and self.llm.ok):
            return rule
        prompt = ("Propose le titre de CV (une seule ligne) qui remplacera le titre actuel pour cette offre.\n"
                  f"- Titre actuel : « {old_title} » ({len(old_title)} caractères) : garde le même style et la même "
                  "ponctuation, et une longueur proche (le titre doit tenir sur la même ligne).\n"
                  "- Reprends l'intitulé du poste de l'offre, éventuellement suivi de 1 ou 2 domaines clés de l'offre.\n"
                  "- N'ajoute ni « senior », ni « manager », ni un niveau ou un diplôme absent du CV.\n"
                  'Réponds en JSON : {"titre": "…"}\n\n'
                  f"OFFRE : {json.dumps({k: offer.get(k) for k in ('poste', 'entreprise', 'competences_requises', 'missions')}, ensure_ascii=False)}\n\n"
                  f"CV (extrait) :\n{cv_text[:3000]}")
        try:
            t = clean_title((self.llm.complete(SYSTEM, prompt, json_mode=True, max_tokens=300) or {}).get("titre", ""))
            return t if 3 < len(t) <= max(len(old_title) + 25, 80) else rule
        except Exception:  # noqa: BLE001
            return rule

    def adapt_cv(self, cv: SourceDoc, old_title: str, new_title: str, offer_text: str, offer: dict) -> Result:
        return self.run_logged(self._adapt_cv, cv, old_title, new_title, offer_text, offer,
                               summary=lambda r: f"titre du CV adapté : « {new_title} »")

    def _adapt_cv(self, cv, old_title, new_title, offer_text, offer) -> Result:
        if not cv.is_docx:
            raise ValueError("Pour garder votre CV strictement identique, importez-le au format Word (.docx) : un PDF ne "
                             "peut pas être modifié sans changer sa mise en page.")
        new_title = new_title.strip()
        if not old_title.strip() or not new_title:
            raise ValueError("Indiquez le titre actuel du CV et le nouveau titre.")
        data, n = replace_text_everywhere(cv.data, old_title, new_title)
        if n == 0:
            raise ValueError(f"Le titre « {old_title} » est introuvable dans le CV.")
        warnings = []
        added = sorted(set(NUM_RX.findall(new_title)) - set(NUM_RX.findall(cv.text)))
        if added:
            warnings.append(f"Le nouveau titre contient un chiffre absent du CV : {', '.join(added)}")
        if len(new_title) > len(old_title) + 15:
            warnings.append("Le nouveau titre est nettement plus long que l'ancien : vérifiez qu'il tient sur une ligne.")
        text = cv.text.replace(old_title, new_title)
        pdf, method = to_pdf(data, text)
        before, after = match_report(cv.text, offer_text), match_report(text, offer_text)
        res = Result("CV", data, pdf, method, text, self._filename("CV", offer),
                     [(old_title, new_title, "Seul changement du CV : le titre, adapté à l'offre")], warnings,
                     before["score"], after["score"], after["missing"])
        # mémorise le titre d'origine pour les prochaines offres
        if (self.cfg.user or {}).get("titre_cv") != old_title and old_title in cv.text:
            try:
                from core.config import save_user_profile
                save_user_profile({**self.cfg.user, "titre_cv": old_title})
                self.cfg.user["titre_cv"] = old_title
            except Exception:  # noqa: BLE001
                pass
        return self._save(res, offer)

    # ------------------------------------------------------------------ lettre
    def letter_layout(self, letter: SourceDoc | None) -> dict:
        """Structure de la lettre modèle : date, objet, formule d'appel, paragraphes de motivation."""
        if not (letter and letter.is_docx):
            return {"date": None, "objet": None, "salut": None, "close": None, "body": []}
        return letter_structure(letter)

    def default_recipient(self, offer: dict) -> list[str]:
        lines = [str(offer.get("entreprise") or "Nom de l'entreprise"), "Service Recrutement"]
        lieu = str(offer.get("lieu") or "").strip()
        lines += ["Adresse", lieu if re.search(r"\d{5}", lieu) else (f"Code postal {lieu}".strip() if lieu else "Code postal Ville")]
        return lines

    def find_address(self, offer: dict) -> tuple[list[str], list[dict]]:
        """Adresse de l'employeur trouvée par recherche web (IA) — à vérifier par l'utilisateur."""
        ent, lieu = str(offer.get("entreprise") or ""), str(offer.get("lieu") or "")
        prompt = (f"Trouve l'adresse postale du siège ou de l'établissement de « {ent} » le plus proche de « {lieu or 'France'} ». "
                  "Réponds UNIQUEMENT par 2 à 3 lignes : numéro et rue, puis code postal et ville. Si tu ne trouves pas, "
                  "réponds « INTROUVABLE ».\n"
                  f"RECHERCHE: {ent} adresse siège {lieu}")
        text, sources = self.llm.research(SYSTEM, prompt, max_tokens=400)
        lines = [re.sub(r"[*`#\[\]]|\(\d+\)", "", l).strip(" -•") for l in text.splitlines() if l.strip()]
        lines = [l for l in lines if l and "introuvable" not in l.lower() and len(l) < 90][:3]
        return ([ent, "Service Recrutement"] + lines) if lines else [], sources

    def propose_motivation(self, letter: SourceDoc | None, cv_text: str, offer_text: str, offer: dict, profile: dict,
                           use_ai: bool = True, instructions: str = "") -> dict:
        lay = self.letter_layout(letter)
        model_paras = [p.text for p in (letter.paragraphs if letter else []) if p.idx in lay["body"]]
        k = len(model_paras) or 3
        poste = clean_title(offer.get("poste")) or "le poste proposé"
        objet = f"Candidature au poste {de(poste)}"
        if use_ai and self.llm.ok:
            lens = [len(t) for t in model_paras] or [450, 500, 350]
            prompt = (
                f"Rédige les {k} paragraphe(s) de motivation d'une lettre de candidature pour cette offre, en remplacement "
                "de ceux de la lettre modèle (le reste de la lettre — coordonnées, formule d'appel, formule de politesse, "
                "signature — ne change pas).\n"
                f"- Exactement {k} paragraphe(s), de longueurs proches de : {lens} caractères (la lettre doit tenir sur une page).\n"
                "- Paragraphe 1 : pourquoi cet employeur et ce poste (élément concret de l'offre ou de l'entreprise).\n"
                "- Ensuite : ce que le candidat apporte, prouvé par des faits RÉELS de son CV, avec le vocabulaire de l'offre.\n"
                "- Dernier : motivation et proposition d'entretien (sans formule de politesse finale).\n"
                "- Garde le ton et la personne grammaticale de la lettre modèle ; aucun chiffre ni fait absent du CV.\n"
                + (f"- Consigne du candidat : {instructions}\n" if instructions.strip() else "")
                + 'Réponds en JSON : {"objet": "Candidature au poste de …", "paragraphes": ["…"]}\n\n'
                f"LETTRE MODÈLE (paragraphes à remplacer) : {json.dumps(model_paras, ensure_ascii=False)}\n\n"
                f"ANALYSE DE L'OFFRE : {json.dumps(offer, ensure_ascii=False)[:2500]}\n\nOFFRE :\n{offer_text[:6000]}\n\n"
                f"CV :\n{cv_text[:8000]}")
            try:
                ai = self.llm.complete(SYSTEM, prompt, json_mode=True, max_tokens=3000, temperature=0.5) or {}
                paras = [str(x).strip() for x in ai.get("paragraphes") or [] if str(x).strip()]
                if paras:
                    warn = sorted(set(NUM_RX.findall(" ".join(paras))) - set(NUM_RX.findall(cv_text + offer_text)))
                    return {"objet": str(ai.get("objet") or objet).replace("Objet :", "").strip(), "paragraphes": paras[:k],
                            "warnings": [f"Chiffres absents du CV et de l'offre à vérifier : {', '.join(warn)}"] if warn else []}
            except Exception:  # noqa: BLE001
                pass
        return {"objet": objet, "paragraphes": self._rule_motivation(k, cv_text, offer, profile), "warnings": []}

    @staticmethod
    def _skill_txt(s: str) -> str:
        s = s.split(" (")[0].replace(" / ", " et ")
        return s if (s[:3].isupper() or re.match(r"^[A-Z]{2,}", s)) else s[:1].lower() + s[1:]

    def _rule_motivation(self, k: int, cv_text: str, offer: dict, profile: dict) -> list[str]:
        poste = clean_title(offer.get("poste")) or "ce poste"
        ent = offer.get("entreprise") or "votre entreprise"
        lieu = f", à {offer['lieu']}" if offer.get("lieu") else ""
        mine = set(find_skills(cv_text))
        asked = offer.get("competences_requises") or []
        matched = [self._skill_txt(s) for s in asked if s in mine and s not in SOFT][:5]
        missions = [str(m).rstrip(".").strip() for m in offer.get("missions") or []][:2]
        titre = (profile.get("titre") or "").strip()
        emps = [e.get("employeur") for e in profile.get("experiences") or [] if isinstance(e, dict) and e.get("employeur")][:2]
        hook = (f"Les missions proposées, notamment {missions[0][:1].lower() + missions[0][1:]}"
                + (f" et {missions[1][:1].lower() + missions[1][1:]}" if len(missions) > 1 else "")
                + ", rejoignent directement mon expérience.") if missions else \
            "Ce poste s'inscrit pleinement dans la continuité de mon parcours et de mon projet professionnel."
        p1 = (f"{'Actuellement ' + titre[:1].lower() + titre[1:] + ', je' if titre else 'Je'} vous adresse ma candidature "
              f"au poste {de(poste)} au sein de {ent}{lieu}. {hook}")
        p2 = ("Au cours de mon parcours" + (f" chez {' puis chez '.join(emps)}" if emps else "")
              + ", j'ai développé des compétences directement mobilisables pour ce poste"
              + (f" : {', '.join(matched)}." if matched else ".")
              + " La rigueur et l'esprit d'analyse que j'ai acquis me permettent de traiter les dossiers avec fiabilité "
                "et de documenter chacune de mes décisions.")
        p3 = (f"Rejoindre {ent} représente pour moi l'opportunité de mettre ces compétences au service de vos équipes et "
              f"de continuer à progresser sur des sujets à fort enjeu. Je serais ravi(e) de vous exposer ma motivation "
              f"lors d'un entretien.")
        base = [p1, p2, p3]
        if k == 1:
            return [" ".join(base)]
        if k == 2:
            return [p1, p2 + " " + p3]
        return base

    def build_letter(self, letter: SourceDoc | None, recipient: list[str], date_line: str, objet: str,
                     paragraphs: list[str], profile: dict, offer: dict, replace_idx: list[int] | None = None) -> Result:
        return self.run_logged(self._build_letter, letter, recipient, date_line, objet, paragraphs, profile, offer,
                               replace_idx, summary=lambda r: f"lettre adaptée ({r.pdf_method})")

    def _build_letter(self, letter, recipient, date_line, objet, paragraphs, profile, offer, replace_idx) -> Result:
        paragraphs = [p.strip() for p in paragraphs if p.strip()]
        recipient = [l.strip() for l in recipient if l.strip()]
        warnings = []
        if letter and letter.is_docx:
            lay = self.letter_layout(letter)
            if not lay["body"]:
                warnings.append("Formule d'appel (« Madame, Monsieur, ») introuvable dans la lettre modèle : la motivation "
                                "n'a pas pu être placée. Vérifiez votre modèle.")
            data = adapt_letter_docx(letter.data, recipient, date_line, objet, paragraphs, replace_idx)
        else:
            if letter:
                warnings.append("Lettre modèle non Word : la lettre est mise en page de façon standard. Importez la version "
                                "Word de votre lettre pour garder sa forme.")
            nom = profile.get("nom") or "Prénom Nom"
            sender = [x for x in (nom, profile.get("localisation"), profile.get("telephone"), profile.get("email")) if x]
            data = standard_letter_docx(sender, recipient, date_line, objet, "Madame, Monsieur,", paragraphs,
                                        "Je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées.",
                                        nom)
        text = read_document(data, "x.docx").text
        pdf, method = to_pdf(data, text, letter=True)
        return self._save(Result("Lettre", data, pdf, method, text, self._filename("LM", offer), [], warnings), offer)

    # ------------------------------------------------------------------ enregistrement
    def _filename(self, prefix: str, offer: dict) -> str:
        nom = safe_filename(self.cfg.user.get("nom") or "candidat")
        poste = safe_filename(clean_title(offer.get("poste")) or "poste")[:50]
        ent = safe_filename(offer.get("entreprise") or "")[:30]
        return "_".join(x for x in (prefix, nom, poste, ent) if x)

    def _save(self, res: Result, offer: dict) -> Result:
        ensure_dirs()
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        pdocx = save_output(OUTPUTS, f"{res.filename}_{stamp}.docx", res.docx)
        ppdf = save_output(OUTPUTS, f"{res.filename}_{stamp}.pdf", res.pdf)
        res.paths = (pdocx, ppdf)
        try:
            self.store.add_document(res.kind, str(offer.get("poste") or ""), str(offer.get("entreprise") or ""),
                                    offer.get("uid"), pdocx, ppdf, {"score_avant": res.score_before, "score_apres": res.score_after})
        except Exception:  # noqa: BLE001
            pass
        return res

    def readiness(self) -> tuple[str, str]:
        from .profil import AgentProfil
        cv, lm = AgentProfil.template_path("cv"), AgentProfil.template_path("lettre")
        if cv and cv.suffix != ".docx":
            return "partiel", "CV modèle en PDF : importez la version Word pour l'adapter sans changer sa forme"
        if cv and lm:
            return "ok", "CV et lettre modèles prêts"
        if cv:
            return "partiel", "CV prêt, pas de lettre modèle (lettre standard)"
        return "manque", "Importez votre CV modèle Word"
