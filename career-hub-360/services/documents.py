"""Documents : lecture (Word, PDF, texte), réécriture d'un modèle Word en gardant sa mise en page, export Word et PDF.

Principe pour le CV et la lettre :
- modèle Word (.docx) → chaque paragraphe modifié est réécrit À SA PLACE (police, taille, couleurs, colonnes, zones de
  texte et tableaux conservés). Les paragraphes contenant un lien hypertexte ne sont pas modifiés.
- modèle PDF ou texte → un nouveau document Word propre est construit à partir du texte adapté.
- PDF : conversion fidèle du Word par LibreOffice s'il est installé (packages.txt sur Streamlit Cloud), sinon PDF
  mis en page par ReportLab.
"""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

FONTS = Path(__file__).resolve().parent.parent / "assets" / "fonts"


# ====================================================================== lecture
@dataclass
class Para:
    idx: int
    text: str
    editable: bool
    style: str = ""


@dataclass
class SourceDoc:
    name: str
    kind: str  # docx | pdf | txt
    data: bytes
    text: str
    paragraphs: list[Para] = field(default_factory=list)

    @property
    def is_docx(self) -> bool:
        return self.kind == "docx"


def _iter_paragraphs(doc) -> list[Paragraph]:
    """Tous les paragraphes du corps, y compris tableaux et zones de texte, dans l'ordre du document."""
    return [Paragraph(p, doc) for p in doc.element.body.iter(qn("w:p"))]


def _in_fallback(p: Paragraph) -> bool:
    el = p._p.getparent()
    while el is not None:
        if el.tag.endswith("}Fallback"):
            return True
        el = el.getparent()
    return False


def _owner_p(el):
    el = el.getparent()
    while el is not None and el.tag != qn("w:p"):
        el = el.getparent()
    return el


def _own(p: Paragraph, tag: str) -> list:
    """Éléments `tag` appartenant à ce paragraphe (pas à une zone de texte imbriquée)."""
    return [e for e in p._p.iter(qn(tag)) if _owner_p(e) is p._p]


def _para_text(p: Paragraph) -> str:
    out = []
    for e in p._p.iter(qn("w:t"), qn("w:br"), qn("w:tab")):
        if _owner_p(e) is not p._p:
            continue
        out.append(e.text or "" if e.tag == qn("w:t") else "\n" if e.tag == qn("w:br") else "\t")
    return "".join(out)


def _has_hyperlink(p: Paragraph) -> bool:
    return any(True for _ in _own(p, "w:hyperlink")) or any(True for _ in _own(p, "w:fldSimple"))


def read_document(data: bytes, name: str) -> SourceDoc:
    ext = Path(name).suffix.lower()
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        paras, seen = [], set()
        for i, p in enumerate(_iter_paragraphs(doc)):
            if _in_fallback(p):  # doublon VML des zones de texte
                continue
            t = _para_text(p)
            editable = bool(t.strip()) and not _has_hyperlink(p)
            paras.append(Para(i, t, editable, p.style.name if p.style is not None else ""))
            seen.add(t)
        text = "\n".join(p.text for p in paras if p.text.strip())
        return SourceDoc(name, "docx", data, text, paras)
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
        text = re.sub(r"[ \t]+", " ", text)
        paras = [Para(i, l, True) for i, l in enumerate(x for x in text.splitlines() if x.strip())]
        return SourceDoc(name, "pdf", data, text.strip(), paras)
    text = data.decode("utf-8", errors="replace")
    paras = [Para(i, l, True) for i, l in enumerate(x for x in text.splitlines() if x.strip())]
    return SourceDoc(name, "txt", data, text.strip(), paras)


# ====================================================================== réécriture d'un modèle Word
def set_paragraph_text(p: Paragraph, new: str) -> None:
    """Remplace le texte d'un paragraphe en gardant la mise en forme.

    Les premiers morceaux (runs) dont le texte est conservé au début du nouveau texte gardent leur style propre
    (ex. libellé en gras « Compétences : ») ; le reste prend le style du morceau suivant. Les morceaux sans texte
    (images, ancres de zones de texte) ne sont jamais touchés.
    """
    runs = [r for r in _own(p, "w:r") if any(True for _ in r.iter(qn("w:t")))
            and not any(True for _ in r.iter(qn("w:drawing"))) and not any(True for _ in r.iter(qn("w:pict")))]
    if not runs:
        p.add_run(new)
        return

    def rtext(r):
        return "".join(t.text or "" for t in r.iter(qn("w:t")))

    def set_rtext(r, text):
        ts = list(r.iter(qn("w:t")))
        for extra in ts[1:]:
            extra.getparent().remove(extra)
        for br in list(r.iter(qn("w:br"))):
            br.getparent().remove(br)
        lines = text.split("\n")
        t0 = ts[0]
        t0.text = lines[0]
        t0.set(qn("xml:space"), "preserve")
        anchor = t0
        for ln in lines[1:]:
            br = r.makeelement(qn("w:br"), {})
            anchor.addnext(br)
            t = r.makeelement(qn("w:t"), {})
            t.text = ln
            t.set(qn("xml:space"), "preserve")
            br.addnext(t)
            anchor = t

    kept, pos = 0, 0
    for r in runs[:-1]:
        t = rtext(r)
        if t and new.startswith(t, pos):
            pos += len(t)
            kept += 1
        else:
            break
    set_rtext(runs[kept], new[pos:])
    for r in runs[kept + 1:]:
        r.getparent().remove(r)


def apply_edits(data: bytes, edits: dict[int, str]) -> bytes:
    """Applique {indice de paragraphe: nouveau texte} à un .docx (les doublons de zones de texte suivent)."""
    doc = Document(io.BytesIO(data))
    paras = _iter_paragraphs(doc)
    by_text = {}
    for i, new in edits.items():
        if 0 <= i < len(paras):
            by_text[_para_text(paras[i])] = new
    for p in paras:
        t = _para_text(p)
        if t in by_text and t.strip() and not _has_hyperlink(p):
            set_paragraph_text(p, by_text[t])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def fill_placeholders(data: bytes, values: dict[str, str]) -> bytes:
    """Remplace {{CLE}} dans un .docx (même si la balise est coupée en plusieurs morceaux)."""
    doc = Document(io.BytesIO(data))
    for p in _iter_paragraphs(doc):
        t = _para_text(p)
        if "{{" in t:
            new = re.sub(r"\{\{\s*([A-Za-zÀ-ÿ_ ]+?)\s*\}\}", lambda m: values.get(m.group(1).strip().upper(), m.group(0)), t)
            if new != t:
                set_paragraph_text(p, new)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ====================================================================== construction d'un Word à partir de texte
_MD_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
_MD_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _add_hyperlink(paragraph, url: str, text: str):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    r = paragraph._p.makeelement(qn("w:r"), {})
    rpr = r.makeelement(qn("w:rPr"), {})
    color = rpr.makeelement(qn("w:color"), {qn("w:val"): "1D4ED8"})
    u = rpr.makeelement(qn("w:u"), {qn("w:val"): "single"})
    rpr.append(color)
    rpr.append(u)
    r.append(rpr)
    t = r.makeelement(qn("w:t"), {})
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def _inline(paragraph, text: str, size: float | None = None):
    """Texte avec **gras** et [liens](url)."""
    pos = 0
    tokens = []
    for m in re.finditer(r"\*\*(.+?)\*\*|\[([^\]]+)\]\((https?://[^)\s]+)\)|(?<![\w*])\*(?!\s)([^*]+?)\*(?![\w*])", text):
        if m.start() > pos:
            tokens.append(("t", text[pos:m.start()]))
        if m.group(1) is not None:
            tokens.append(("b", m.group(1)))
        elif m.group(4) is not None:
            tokens.append(("i", m.group(4)))
        else:
            tokens.append(("l", m.group(2), m.group(3)))
        pos = m.end()
    tokens.append(("t", text[pos:]))
    for tk in tokens:
        if tk[0] == "l":
            _add_hyperlink(paragraph, tk[2], tk[1])
            continue
        if not tk[1]:
            continue
        run = paragraph.add_run(tk[1])
        run.bold = tk[0] == "b"
        run.italic = tk[0] == "i"
        if size:
            run.font.size = Pt(size)


def normalize_md(md: str) -> str:
    """Simplifie le Markdown des IA : titres profonds → ###, tableaux → puces, sources [n] conservées."""
    out = []
    for line in str(md or "").splitlines():
        s = line.strip()
        if re.match(r"^\|?\s*:?-{3,}", s):
            continue
        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|") if c.strip()]
            out.append("- " + " · ".join(cells))
            continue
        m = re.match(r"^(#{4,6})\s+(.*)$", s)
        out.append(f"### {m.group(2)}" if m else line)
    return "\n".join(out)


def markdown_to_docx(md: str, base_font: str = "Calibri", accent: str = "1F3A5F") -> bytes:
    """Markdown simple (#, ##, ###, -, **gras**, *italique*, [lien](url), ---) → .docx propre."""
    md = normalize_md(md)
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = base_font
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), base_font)
    rgb = RGBColor.from_string(accent)
    for line in str(md or "").splitlines():
        raw = line.rstrip()
        s = raw.strip()
        if not s:
            continue
        if s in ("---", "***"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run("_" * 40).font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            run = p.add_run(_MD_BOLD.sub(r"\1", _MD_LINK.sub(r"\1", m.group(2))))
            run.bold = True
            run.font.size = Pt({1: 17, 2: 13, 3: 11.5}[level])
            run.font.color.rgb = rgb
            p.paragraph_format.space_before = Pt({1: 0, 2: 10, 3: 6}[level])
            p.paragraph_format.space_after = Pt(3)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue
        m = re.match(r"^[-*•]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, m.group(1))
            p.paragraph_format.space_after = Pt(1)
            continue
        p = doc.add_paragraph()
        _inline(p, s)
        p.paragraph_format.space_after = Pt(4)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def plain_letter_docx(lines: list[str], base_font: str = "Calibri") -> bytes:
    """Lettre : un paragraphe par élément, sans mise en forme de titre."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)
    st = doc.styles["Normal"]
    st.font.name = base_font
    st.font.size = Pt(11)
    for l in lines:
        p = doc.add_paragraph()
        _inline(p, l)
        p.paragraph_format.space_after = Pt(8)
        if l.lower().startswith("objet"):
            for r in p.runs:
                r.bold = True
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def docx_text(data: bytes) -> str:
    return read_document(data, "x.docx").text


# ====================================================================== PDF
def soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def docx_to_pdf(data: bytes, timeout: int = 120) -> bytes | None:
    """Conversion fidèle Word → PDF par LibreOffice (None si indisponible ou en échec)."""
    exe = soffice_path()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "document.docx"
        src.write_bytes(data)
        profile = Path(d) / "lo-profile"
        try:
            subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                            "--outdir", d, str(src)], capture_output=True, timeout=timeout, check=False)
        except (subprocess.SubprocessError, OSError):
            return None
        pdf = Path(d) / "document.pdf"
        return pdf.read_bytes() if pdf.exists() and pdf.stat().st_size > 0 else None


def _register_fonts() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVu", str(FONTS / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONTS / "DejaVuSans-Bold.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVu-Oblique", str(FONTS / "DejaVuSans-Oblique.ttf")))
            from reportlab.lib.fonts import addMapping
            addMapping("DejaVu", 0, 0, "DejaVu")
            addMapping("DejaVu", 1, 0, "DejaVu-Bold")
            addMapping("DejaVu", 0, 1, "DejaVu-Oblique")
            addMapping("DejaVu", 1, 1, "DejaVu-Bold")
        return "DejaVu"
    except Exception:  # noqa: BLE001
        return "Helvetica"


def _rl_inline(text: str) -> str:
    from xml.sax.saxutils import escape
    out, pos = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\[([^\]]+)\]\((https?://[^)\s]+)\)|(?<![\w*])\*(?!\s)([^*]+?)\*(?![\w*])", text):
        out.append(escape(text[pos:m.start()]))
        if m.group(1) is not None:
            out.append(f"<b>{escape(m.group(1))}</b>")
        elif m.group(4) is not None:
            out.append(f"<i>{escape(m.group(4))}</i>")
        else:
            out.append(f'<link href="{escape(m.group(3))}" color="#1D4ED8"><u>{escape(m.group(2))}</u></link>')
        pos = m.end()
    out.append(escape(text[pos:]))
    return "".join(out)


def markdown_to_pdf(md: str, letter: bool = False) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph as P, SimpleDocTemplate, Spacer

    md = normalize_md(md)
    font = _register_fonts()
    base = ParagraphStyle("b", fontName=font, fontSize=10.5 if not letter else 11, leading=14.5 if not letter else 16,
                          spaceAfter=4 if not letter else 9)
    h = {1: ParagraphStyle("h1", parent=base, fontSize=17, leading=21, textColor=colors.HexColor("#1F3A5F"), spaceAfter=4,
                           fontName=font + "-Bold" if font == "DejaVu" else "Helvetica-Bold"),
         2: ParagraphStyle("h2", parent=base, fontSize=12.5, leading=16, textColor=colors.HexColor("#1F3A5F"), spaceBefore=8,
                           fontName=font + "-Bold" if font == "DejaVu" else "Helvetica-Bold"),
         3: ParagraphStyle("h3", parent=base, fontSize=11, leading=14, spaceBefore=4,
                           fontName=font + "-Bold" if font == "DejaVu" else "Helvetica-Bold")}
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    flow, bullets = [], []

    def flush():
        if bullets:
            flow.append(ListFlowable([ListItem(P(_rl_inline(b), base), leftIndent=12) for b in bullets],
                                     bulletType="bullet", start="•", leftIndent=12, bulletFontName=font))
            bullets.clear()

    for line in str(md or "").splitlines():
        s = line.strip()
        if not s:
            flush()
            continue
        if s in ("---", "***"):
            flush()
            flow.append(HRFlowable(width="100%", color=colors.HexColor("#CBD5E1"), spaceBefore=3, spaceAfter=5))
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            flush()
            flow.append(P(_rl_inline(m.group(2)), h[len(m.group(1))]))
            continue
        m = re.match(r"^[-*•]\s+(.*)$", s)
        if m:
            bullets.append(m.group(1))
            continue
        flush()
        flow.append(P(_rl_inline(s), base))
    flush()
    if not flow:
        flow.append(Spacer(1, 1))
    doc.build(flow)
    return buf.getvalue()


def to_pdf(docx_bytes: bytes | None, md_fallback: str, letter: bool = False) -> tuple[bytes, str]:
    """PDF fidèle au Word si LibreOffice est présent ; sinon PDF construit à partir du texte. Renvoie (pdf, méthode)."""
    if docx_bytes:
        pdf = docx_to_pdf(docx_bytes)
        if pdf:
            return pdf, "LibreOffice (mise en page du Word conservée)"
    return markdown_to_pdf(md_fallback, letter=letter), "ReportLab (mise en page simplifiée)"


def safe_filename(s: str, default: str = "document") -> str:
    s = re.sub(r"[^\w\-]+", "_", str(s or "").strip(), flags=re.UNICODE).strip("_")
    return (s or default)[:80]


def save_output(folder: Path, name: str, data: bytes) -> str:
    folder.mkdir(parents=True, exist_ok=True)
    p = folder / name
    p.write_bytes(data)
    return str(p)


def file_exists(p: str | None) -> bool:
    return bool(p) and os.path.exists(p)
